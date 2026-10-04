"""
REQ-3: Writes the 3-sheet Excel workbook from the validator's processed output.
  Sheet 1: Extracted_Attributes - high-confidence, explicitly-stated values
  Sheet 2: Needs_Review         - low-confidence / inferred / schema-flagged values
  Sheet 3: Quality_Summary      - classification, completeness, missing critical
                                   attributes, quality rule notes
"""

import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from extraction_agent import run_extraction
from validator import load_dictionary, validate_and_process

HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
HEADER_FONT = Font(bold=True)


def _write_header(ws, headers):
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    ws.freeze_panes = "A2"


def _autosize(ws, max_width=60):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 2, max_width)


def write_workbook(processed: dict, pdf_path: str, output_path: str):
    wb = Workbook()

    # --- Sheet 1: Extracted_Attributes ---
    ws1 = wb.active
    ws1.title = "Extracted_Attributes"
    headers1 = ["Attribute Key", "Value", "Unit", "Source Document", "Source Quote", "Confidence", "Conflict Flag"]
    _write_header(ws1, headers1)
    for row in processed["extracted_attributes"]:
        ws1.append([
            row["key"],
            row["value"],
            row.get("unit") or "",
            os.path.basename(pdf_path),
            row.get("source_quote") or "",
            row.get("confidence"),
            "",  # single-document runs rarely produce conflicts; see ADR note
        ])
    _autosize(ws1)

    # --- Sheet 2: Needs_Review ---
    ws2 = wb.create_sheet("Needs_Review")
    headers2 = ["Attribute Key", "Value", "Unit", "Confidence", "Reason Flagged", "Source Document", "Source Quote"]
    _write_header(ws2, headers2)
    for row in processed["needs_review"]:
        ws2.append([
            row["key"],
            row.get("value"),
            row.get("unit") or "",
            row.get("confidence"),
            row["reason"],
            os.path.basename(pdf_path),
            row.get("source_quote") or "",
        ])
    _autosize(ws2)

    # --- Sheet 3: Quality_Summary ---
    ws3 = wb.create_sheet("Quality_Summary")
    ws3.append(["Field", "Value"])
    ws3["A1"].font = HEADER_FONT
    ws3["B1"].font = HEADER_FONT
    ws3["A1"].fill = HEADER_FILL
    ws3["B1"].fill = HEADER_FILL

    classification = processed["classification"]
    ws3.append(["Source Document", os.path.basename(pdf_path)])
    ws3.append(["Ingredient Type Classification", classification.get("ingredient_type")])
    ws3.append(["Classification Evidence", classification.get("evidence_quote")])
    ws3.append(["Completeness Score (critical attributes)", f"{processed['completeness_score']}%"])
    ws3.append(["Missing Critical Attributes", ", ".join(processed["missing_critical_attributes"]) or "None"])
    ws3.append(["", ""])
    ws3.append(["Quality Rule Notes", ""])
    for note in processed["quality_notes"]:
        ws3.append([note["rule_id"], note["note"]])
    if not processed["quality_notes"]:
        ws3.append(["", "No quality rule violations detected."])
    _autosize(ws3)

    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python excel_writer.py <path_to_pdf>")
        sys.exit(1)

    pdf_path = sys.argv[1]
    dictionary = load_dictionary()
    raw_result = run_extraction(pdf_path)
    processed = validate_and_process(raw_result, dictionary)

    base_name = os.path.splitext(os.path.basename(pdf_path))[0]
    os.makedirs("output", exist_ok=True)
    output_path = f"output/{base_name}_extraction.xlsx"

    write_workbook(processed, pdf_path, output_path)
    print(f"Wrote {output_path}")
    print(f"  Extracted_Attributes: {len(processed['extracted_attributes'])} rows")
    print(f"  Needs_Review: {len(processed['needs_review'])} rows")
    print(f"  Completeness: {processed['completeness_score']}%")